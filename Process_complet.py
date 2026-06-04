import cv2
import numpy as np
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl_image_loader import SheetImageLoader
from openpyxl.utils import get_column_letter  # Nouvel outil pour convertir un numéro en lettre (ex: 5 -> E)

print("=" * 80 + "\nANALYSE POLYMERE -> PIPELINE MACHINE LEARNING (100% AUTO)\n" + "=" * 80)

# ---------------------------------------------------------
# 1. CHARGEMENT DU FICHIER EXCEL SOUCHE
# ---------------------------------------------------------
fichier_excel = "Electrospinning_experiments.xlsx"

if not os.path.exists(fichier_excel):
    print(f"[ERREUR FATALE] Le fichier {fichier_excel} est introuvable.")
    exit()

# Chargement pour les données avec Pandas
df = pd.read_excel(fichier_excel)
df['Score Global'] = None
df['Classe'] = None

# --- RECHERCHE AUTOMATIQUE DE LA COLONNE "Picture" ---
if 'Picture' not in df.columns:
    print("[ERREUR FATALE] La colonne 'Picture' est introuvable dans l'Excel. Vérifie l'orthographe.")
    exit()

# On trouve à quelle position se trouve 'Picture' (Pandas commence à 0, Excel à 1, d'où le +1)
index_colonne_picture = df.columns.get_loc('Picture') + 1
lettre_colonne = get_column_letter(index_colonne_picture)
print(f"-> Colonne 'Picture' détectée automatiquement dans la colonne : {lettre_colonne}")

# Chargement spécifique pour les IMAGES avec Openpyxl
print("Chargement des images depuis Excel en mémoire (cela peut prendre quelques secondes)...")
wb = load_workbook(fichier_excel)
sheet = wb.active
image_loader = SheetImageLoader(sheet)

# ---------------------------------------------------------
# 2. BOUCLE D'ANALYSE DYNAMIQUE
# ---------------------------------------------------------
# iterrows() permet de parcourir TOUTES les lignes, peu importe combien il y en a !
for index_excel, row in df.iterrows():
    # La ligne 1 dans Excel, ce sont les titres. Donc l'index 0 de Pandas = ligne 2 d'Excel.
    ligne_excel = index_excel + 2

    # Nom de cellule 100% dynamique (ex: "E2", "F3", etc.)
    nom_cellule = f"{lettre_colonne}{ligne_excel}"

    try:
        pil_image = image_loader.get(nom_cellule)
        img_array = np.array(pil_image)
        if len(img_array.shape) == 3 and img_array.shape[2] >= 3:
            img = cv2.cvtColor(img_array, cv2.COLOR_RGB2BGR)
        else:
            img = cv2.cvtColor(img_array, cv2.COLOR_GRAY2BGR)

    except Exception as e:
        print(f"Ligne {ligne_excel} (Cellule {nom_cellule}) : IMAGE INTROUVABLE ou vide.")
        continue  # On passe à la ligne suivante

    # --- TRAITEMENT D'IMAGE ---
    new_w = 500
    new_h = int(img.shape[0] * (new_w / img.shape[1]))
    img_resized = cv2.resize(img, (new_w, new_h))

    hsv = cv2.cvtColor(img_resized, cv2.COLOR_BGR2HSV)
    gray = cv2.cvtColor(img_resized, cv2.COLOR_BGR2GRAY)

    # Masque Carton
    masque_carton = cv2.inRange(hsv, np.array([5, 40, 50]), np.array([40, 255, 255]))
    masque_carton = cv2.morphologyEx(masque_carton, cv2.MORPH_CLOSE, np.ones((5, 5), np.uint8))
    contours_carton, _ = cv2.findContours(masque_carton, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    x1, y1, x2, y2 = 50, 50, new_w - 50, new_h - 50
    if contours_carton:
        x, y, w, h = cv2.boundingRect(max(contours_carton, key=cv2.contourArea))
        marge = 10
        x1, y1, x2, y2 = x + marge, y + marge, x + w - marge, y + h - marge

    zone_utile = np.zeros_like(masque_carton)
    cv2.rectangle(zone_utile, (x1, y1), (x2, y2), 255, -1)

    # Masque Polymère
    masque_brut = cv2.inRange(hsv, np.array([0, 0, 170]), np.array([180, 55, 255]))
    masque_filtre = cv2.bitwise_and(masque_brut, zone_utile)
    contours, _ = cv2.findContours(masque_filtre, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    contours_valides = [c for c in contours if cv2.contourArea(c) > 500]
    nb_depots = len(contours_valides)

    if nb_depots == 0:
        print(f"Ligne {ligne_excel} : Aucun dépôt détecté.")
        continue

    plus_grand_contour = max(contours_valides, key=cv2.contourArea)
    masque_propre = np.zeros_like(masque_filtre)
    cv2.drawContours(masque_propre, [plus_grand_contour], -1, 255, -1)

    # Analyse Géométrique
    dist_transform = cv2.distanceTransform(masque_propre, cv2.DIST_L2, 5)
    _, max_val, _, max_loc = cv2.minMaxLoc(dist_transform)
    centre_x, centre_y = max_loc
    rayon_coeur = int(max_val)

    dist_max_haut, dist_max_bas = 0, 0
    for point in plus_grand_contour:
        px, py = point[0]
        dist = np.sqrt((px - centre_x) ** 2 + (py - centre_y) ** 2)
        if py < centre_y and dist > dist_max_haut:
            dist_max_haut = dist
        elif py >= centre_y and dist > dist_max_bas:
            dist_max_bas = dist

    ratio_haut, ratio_bas = dist_max_haut / rayon_coeur, dist_max_bas / rayon_coeur
    score_geo = max(0, (max(ratio_haut, ratio_bas) - 1) * 100)

    # Analyse Texture
    masque_coeur = np.zeros_like(gray)
    cv2.circle(masque_coeur, (centre_x, centre_y), rayon_coeur, 255, -1)
    details = cv2.absdiff(gray, cv2.medianBlur(gray, 21))
    pixels = details[masque_coeur == 255]

    score_texture = 0
    if len(pixels) > 0:
        ratio_leger = (np.sum((pixels > 10) & (pixels <= 22)) / len(pixels)) * 100
        ratio_critique = (np.sum(pixels > 22) / len(pixels)) * 100
        malus_intensite = max(0, np.percentile(pixels, 99.5) - 22) * 3
        score_texture = ratio_leger * 0.2 + ratio_critique * 300 + malus_intensite

    # Calcul Score Global
    score_global = (score_geo * 0.60) + (score_texture * 0.40)

    if nb_depots >= 2:
        num_classe = 4
    else:
        num_classe = 1 if score_global <= 30 else 2 if score_global <= 45 else 3 if score_global <= 75 else 4

    df.loc[index_excel, 'Score Global'] = round(score_global, 2)
    df.loc[index_excel, 'Classe'] = num_classe
    print(f"Ligne {ligne_excel} analysée -> Label : {num_classe}")

# ---------------------------------------------------------
# 3. SAUVEGARDE EN CSV
# ---------------------------------------------------------
fichier_sortie = "results.csv"
# Optionnel : Si tu ne veux pas que la colonne 'Picture' (qui ne contient plus rien d'utile en texte) pollue ton CSV
# df = df.drop(columns=['Picture'])

df.to_csv(fichier_sortie, index=False, sep=",")

print("\n" + "=" * 80)
print(f"FICHIER IA PRÊT ! Sauvegardé sous : {fichier_sortie}")
print("=" * 80)